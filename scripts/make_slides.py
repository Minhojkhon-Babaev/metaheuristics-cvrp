"""Сборка презентации по шаблону МИЭМ.

Слайды не рисуются с нуля: эталонные слайды шаблона клонируются вместе с их
геометрией, шрифтами и цветовой системой, после чего в них подставляется текст.
Так сохраняется визуальный стиль оригинала.
"""

from __future__ import annotations

import argparse
import csv
import statistics
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path.home() / ".codex/skills/artifact-template-babaev-minhojkhon-miem/assets/reference.pptx"

# Индексы эталонных слайдов, используемых как макеты.
SRC_TITLE = 0
SRC_THREE_CARDS = 1  # два блока сверху + широкий блок + подпись снизу
SRC_FOUR_CARDS = 5  # четыре карточки с акцентной полосой
SRC_SIX_STEPS = 7  # шесть пронумерованных шагов


# ----------------------------------------------------------- работа с pptx


def clone_slide(prs: Presentation, source):
    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in source.shapes:
        slide.shapes._spTree.append(deepcopy(shape._element))
    return slide


def delete_slides(prs: Presentation, count: int) -> None:
    id_list = prs.slides._sldIdLst
    for element in list(id_list)[:count]:
        prs.part.drop_rel(element.rId)
        id_list.remove(element)


def by_name(slide) -> dict:
    return {shape.name: shape for shape in slide.shapes}


def set_text(shape, lines) -> None:
    """Переписывает текст фигуры, сохраняя форматирование первого абзаца."""
    if isinstance(lines, str):
        lines = [lines]
    body = shape.text_frame._txBody
    paragraphs = body.findall(qn("a:p"))
    template = None
    for paragraph in paragraphs:
        if paragraph.findall(qn("a:r")):
            template = deepcopy(paragraph)
            break
    if template is None:
        template = deepcopy(paragraphs[0])

    for paragraph in paragraphs:
        body.remove(paragraph)

    for line in lines:
        paragraph = deepcopy(template)
        for br in paragraph.findall(qn("a:br")):
            paragraph.remove(br)
        runs = paragraph.findall(qn("a:r"))
        for extra in runs[1:]:
            paragraph.remove(extra)
        if runs:
            runs[0].find(qn("a:t")).text = line
        body.append(paragraph)


def keep_only(slide, names: set[str]) -> None:
    for shape in list(slide.shapes):
        if shape.name not in names:
            shape._element.getparent().remove(shape._element)


def add_picture_slide(prs: Presentation, source, title: str, title_name: str, image: Path, top=1.15):
    slide = clone_slide(prs, source)
    keep_only(slide, {title_name})
    set_text(by_name(slide)[title_name], title)
    width = Inches(11.6)
    picture = slide.shapes.add_picture(str(image), Inches(0.0), Inches(top), width=width)
    picture.left = int((prs.slide_width - picture.width) / 2)
    max_height = Inches(7.5 - top - 0.25)
    if picture.height > max_height:
        scale = max_height / picture.height
        picture.height = int(picture.height * scale)
        picture.width = int(picture.width * scale)
        picture.left = int((prs.slide_width - picture.width) / 2)
    return slide


# ------------------------------------------------------------------ данные


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def collect_facts(results: Path) -> dict:
    instances = read_csv(results / "summary_per_instance.csv")
    by_set = {row["group"]: row for row in read_csv(results / "summary_by_set.csv")}
    by_size = read_csv(results / "summary_by_size.csv")

    mean_gap = statistics.mean(float(r["mean_gap"]) for r in instances)
    best_gap = statistics.mean(float(r["best_gap"]) for r in instances)
    init_gap = statistics.mean(float(r["init_gap"]) for r in instances)
    solved = sum(int(r["solved_to_optimum"]) for r in instances)
    worst = max(instances, key=lambda r: float(r["mean_gap"]))

    budget_rows = []
    path = results / "tuning_budget.csv"
    if path.exists():
        grouped = defaultdict(list)
        for row in read_csv(path):
            grouped[(row["param"], row["value"])].append(float(row["gap"]))
        budget_rows = [(k[0], k[1], statistics.mean(v)) for k, v in grouped.items()]

    return {
        "instances": instances,
        "by_set": by_set,
        "by_size": by_size,
        "mean_gap": mean_gap,
        "best_gap": best_gap,
        "init_gap": init_gap,
        "solved": solved,
        "total": len(instances),
        "worst": worst,
        "budget_rows": budget_rows,
    }


# ------------------------------------------------------------------ слайды


def build(prs: Presentation, facts: dict, figures: Path, repo_url: str) -> None:
    src = list(prs.slides)
    original = len(src)

    # 1. Титул
    slide = clone_slide(prs, src[SRC_TITLE])
    shapes = by_name(slide)
    names = list(shapes)
    set_text(shapes[names[0]], ["Московский институт электроники", "и математики имени А.Н. Тихонова"])
    set_text(shapes[names[1]], ["Методы оптимизации:", "метаэвристики"])
    set_text(shapes[names[2]], ["Москва", "2026"])
    set_text(
        shapes[names[3]],
        [
            "Метаэвристика ALNS для задачи маршрутизации",
            "транспорта с ограничениями грузоподъёмности:",
            "эксперимент на наборах E, F, M, P (CVRPLIB)",
        ],
    )
    set_text(shapes[names[4]], ["Выполнил:", "Бабаев Минходжхон Зафарович"])

    # 2. Задача, данные, результат
    slide = clone_slide(prs, src[SRC_THREE_CARDS])
    shapes = by_name(slide)
    names = sorted(shapes)
    set_text(shapes["Google Shape;85;p17"], "Задача, данные и метрика")
    set_text(shapes["Google Shape;88;p17"], "Задача (CVRP)")
    set_text(
        shapes["Google Shape;89;p17"],
        [
            "Депо, n клиентов со спросом, однотипные машины вместимости Q",
            "Каждый клиент обслужен ровно один раз, загрузка маршрута ≤ Q",
            "Минимизируется суммарная длина; число машин равно k из имени инстанса",
            "NP-трудная задача — нужен метаэвристический подход",
        ],
    )
    set_text(shapes["Google Shape;92;p17"], "Данные: CVRPLIB")
    set_text(
        shapes["Google Shape;93;p17"],
        [
            f"Все {facts['total']} инстансов наборов E, F, M, P; n от 12 до 199",
            "E — Christofides & Eilon, F — реальные данные Fisher,",
            "M — крупные задачи, P — варьируемый размер парка",
            "Метрика EUC_2D с округлением до целого + один EXPLICIT-инстанс",
        ],
    )
    set_text(shapes["Google Shape;96;p17"], "Контроль корректности метрики")
    set_text(
        shapes["Google Shape;97;p17"],
        [
            "Стоимость эталонных решений .sol пересчитана нашей функцией расстояния "
            f"и совпала с опубликованными оптимумами на всех {facts['total']} инстансах.",
            "Каждое выданное решение проходит строгую валидацию: покрытие всех клиентов, "
            "отсутствие повторов, соблюдение вместимости, сходимость кэшированной стоимости.",
        ],
    )
    set_text(
        shapes["Google Shape;98;p17"],
        f"Итог: среднее отклонение от оптимума {facts['mean_gap']:.2f}% "
        f"при допустимых 10% — точный оптимум найден на {facts['solved']} из {facts['total']} инстансов",
    )

    # 3. Схема алгоритма
    slide = clone_slide(prs, src[SRC_SIX_STEPS])
    shapes = by_name(slide)
    set_text(shapes["Google Shape;260;p23"], "Алгоритм: ALNS + гранулярный локальный поиск")
    steps = [
        (
            "Google Shape;264;p23",
            "Google Shape;265;p23",
            "Стартовое решение",
            "Параллельный алгоритм сбережений Кларка — Райта; для разных сидов сбережения "
            "зашумляются, что даёт разные точки старта",
        ),
        (
            "Google Shape;269;p23",
            "Google Shape;270;p23",
            "Разрушение",
            "5 операторов: random, worst, Shaw (похожие клиенты), string (SISR-отрезки), "
            "route (расформирование слабо загруженных маршрутов)",
        ),
        (
            "Google Shape;274;p23",
            "Google Shape;275;p23",
            "Восстановление",
            "Жадная вставка в 3 порядках обхода и regret-2; позиции ищутся только рядом "
            "с K ближайшими соседями, вставка с «морганием»",
        ),
        (
            "Google Shape;279;p23",
            "Google Shape;280;p23",
            "Локальный поиск",
            "relocate, swap, 2-opt, Or-opt (сегменты 2–3), межмаршрутный 2-opt*; "
            "обход по очереди «грязных» вершин, соседства ограничены K",
        ),
        (
            "Google Shape;284;p23",
            "Google Shape;285;p23",
            "Критерий приёма",
            "Имитация отжига: стартовая температура по схеме Ропке — Писингера, "
            "охлаждение привязано к доле израсходованного бюджета",
        ),
        (
            "Google Shape;289;p23",
            "Google Shape;290;p23",
            "Адаптация и остановка",
            "Веса операторов пересчитываются раз в 100 итераций по их наградам; "
            "стоп — по стагнации рекорда либо по бюджету времени T(n)",
        ),
    ]
    for title_name, body_name, title, body in steps:
        set_text(shapes[title_name], title)
        set_text(shapes[body_name], body)

    # 4. Подбор параметров и критерий остановки
    slide = clone_slide(prs, src[SRC_FOUR_CARDS])
    shapes = by_name(slide)
    set_text(shapes["Google Shape;214;p21"], "Подбор параметров и критерий остановки")
    set_text(shapes["Google Shape;217;p21"], "Методика")
    set_text(
        shapes["Google Shape;218;p21"],
        "OFAT на отдельных 9 настроечных инстансах × 5 сидов, бюджет 10 с (900 прогонов). "
        "Сравнение парное — по одним и тем же парам «инстанс–сид», что снимает разброс между задачами.",
    )
    set_text(shapes["Google Shape;221;p21"], "Значимы только два фактора")
    set_text(
        shapes["Google Shape;222;p21"],
        "Отказ от отжига (T = 0) даёт +0.14 п.п., разрушение до 70% клиентов +0.10 п.п. "
        "К размеру соседства K, адаптации весов и «морганию» алгоритм устойчив: эффекты в пределах 1σ.",
    )
    set_text(shapes["Google Shape;225;p21"], "Контрольный прогон")
    set_text(
        shapes["Google Shape;226;p21"],
        "Объединение всех «лучших» значений OFAT оказалось хуже базовой на +0.09 ± 0.03 п.п. — "
        "параметры взаимодействуют. Итог: K = 20, удаление 10–40%, T₀ = 3% стоимости, λ = 0.3, blink = 0.01.",
    )
    set_text(shapes["Google Shape;228;p21"], "OFAT-победители не складываются — нужна проверка комбинации")
    set_text(shapes["Google Shape;231;p21"], "Критерий остановки")
    set_text(
        shapes["Google Shape;232;p21"],
        "Первое из двух: 30 000 итераций без нового рекорда либо бюджет T(n) = clip(0.3·n, 15, 90) с. "
        "Кривая «качество — время» насыщается после ~10 с; малые инстансы стопятся по стагнации за секунды.",
    )
    set_text(shapes["Google Shape;234;p21"], "Стагнация как основной критерий, время — как страховка")

    add_picture_slide(
        prs,
        src[SRC_FOUR_CARDS],
        "Результаты: отклонение по типам задач",
        "Google Shape;214;p21",
        figures / "results_summary.png",
    )
    add_picture_slide(
        prs,
        src[SRC_FOUR_CARDS],
        "Зависимость качества и скорости от размерности",
        "Google Shape;214;p21",
        figures / "scaling.png",
    )

    # 7. Выводы
    slide = clone_slide(prs, src[SRC_THREE_CARDS])
    shapes = by_name(slide)
    set_text(shapes["Google Shape;85;p17"], "Выводы")
    set_text(shapes["Google Shape;88;p17"], "Качество")
    worst = facts["worst"]
    set_text(
        shapes["Google Shape;89;p17"],
        [
            f"Среднее отклонение {facts['mean_gap']:.2f}% — втрое с запасом внутри порога 10%",
            f"Лучший из 5 запусков: {facts['best_gap']:.2f}%",
            f"Оптимум найден точно на {facts['solved']} из {facts['total']} инстансов",
            f"Худший инстанс — {worst['instance']} ({float(worst['mean_gap']):.2f}%)",
        ],
    )
    set_text(shapes["Google Shape;92;p17"], "Зависимость от типа и размера")
    set_text(
        shapes["Google Shape;93;p17"],
        [
            "E, F, P решаются практически точно при любом n",
            "Набор M — единственный проблемный: плотная упаковка (запас вместимости <1%)",
            "Отклонение растёт с n, время до рекорда — быстрее, чем линейно",
            "Скорость: число итераций в секунду падает примерно как n^-1",
        ],
    )
    set_text(shapes["Google Shape;96;p17"], "Что определило результат")
    set_text(
        shapes["Google Shape;97;p17"],
        [
            "Гранулярность (K ближайших соседей) в локальном поиске и во вставке — без неё "
            "на n ≈ 200 не хватило бы итераций; отжиг вместо чистого спуска; "
            "оператор расформирования маршрутов, позволяющий сократить число машин.",
        ],
    )
    set_text(shapes["Google Shape;98;p17"], f"Репозиторий с кодом и результатами: {repo_url}")

    delete_slides(prs, original)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", default=str(ROOT / "results"))
    parser.add_argument("--out", default=str(ROOT / "slides" / "cvrp_alns.pptx"))
    parser.add_argument("--repo", default="https://github.com/Minhojkhon-Babaev/metaheuristics-cvrp")
    args = parser.parse_args()

    results = Path(args.results)
    prs = Presentation(str(TEMPLATE))
    build(prs, collect_facts(results), results / "figures", args.repo)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"слайдов: {len(prs.slides._sldIdLst)}; сохранено: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
